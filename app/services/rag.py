from __future__ import annotations

import hashlib
import math
import re
import uuid
from pathlib import Path
from typing import Iterable

from sqlalchemy.orm import Session

from app.config import settings
from app.models import KnowledgeDocument, KnowledgeChunk

BASE = Path(__file__).resolve().parents[2]
DOC_DIR = BASE / "data" / "documents"
DOC_DIR.mkdir(parents=True, exist_ok=True)

ALLOWED = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "text/plain": ".txt",
}

# The embedding model is loaded lazily. Cloud deployments can sometimes be
# unable to download a Hugging Face model on the first request, so RAG has a
# deterministic local fallback that requires no model download.
_embedder = None
_embedding_mode: str | None = None
_embedding_error: str | None = None
_client = None
FALLBACK_DIMENSIONS = 384


def _load_dependencies():
    global _client

    if _client is None:
        import chromadb

        _client = chromadb.PersistentClient(
            path=str(BASE / "data" / "vector_store")
        )

    return _client


def _load_embedder():
    global _embedder, _embedding_mode, _embedding_error

    if _embedding_mode == "hash":
        return None

    if _embedder is not None:
        return _embedder

    try:
        from sentence_transformers import SentenceTransformer

        _embedder = SentenceTransformer(settings.embedding_model)
        _embedding_mode = "sentence-transformer"
        _embedding_error = None
        return _embedder
    except Exception as exc:
        # Do not make document upload fail just because a remote model cannot
        # be downloaded/loaded in the deployment environment.
        _embedding_mode = "hash"
        _embedding_error = f"{type(exc).__name__}: {exc}"
        return None


def _fallback_embedding(text: str) -> list[float]:
    """Create a deterministic, normalized local text vector.

    This is intentionally model-free. It is used only when the transformer
    model cannot be loaded, keeping RAG usable in restricted cloud runtimes.
    """
    vector = [0.0] * FALLBACK_DIMENSIONS
    normalized = re.sub(r"\s+", " ", text.lower()).strip()
    tokens = re.findall(r"[a-z0-9]+", normalized)

    features = list(tokens)
    features.extend(
        f"{tokens[i]} {tokens[i + 1]}"
        for i in range(len(tokens) - 1)
    )

    for feature in features:
        digest = hashlib.blake2b(
            feature.encode("utf-8"),
            digest_size=8,
        ).digest()
        index = int.from_bytes(digest, "big") % FALLBACK_DIMENSIONS
        vector[index] += 1.0

    # Character trigrams make the fallback a little more tolerant of
    # spelling/word-form differences.
    compact = re.sub(r"\s+", " ", normalized)
    for i in range(max(0, len(compact) - 2)):
        feature = compact[i : i + 3]
        digest = hashlib.blake2b(
            ("char:" + feature).encode("utf-8"),
            digest_size=8,
        ).digest()
        index = int.from_bytes(digest, "big") % FALLBACK_DIMENSIONS
        vector[index] += 0.15

    norm = math.sqrt(sum(x * x for x in vector))
    if norm == 0:
        vector[0] = 1.0
        norm = 1.0

    return [x / norm for x in vector]


def _embed_texts(texts: list[str]) -> list[list[float]]:
    embedder = _load_embedder()

    if embedder is not None:
        try:
            return embedder.encode(
                texts,
                normalize_embeddings=True,
            ).tolist()
        except Exception as exc:
            global _embedding_mode, _embedding_error
            _embedding_mode = "hash"
            _embedding_error = f"{type(exc).__name__}: {exc}"

    return [_fallback_embedding(text) for text in texts]


def _collection_name(business_id: int) -> str:
    return f"{settings.rag_collection_prefix}_{business_id}"


def _extract_pdf(path: Path) -> list[tuple[str, int | None, str | None]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return [
        (page.extract_text() or "", i + 1, None)
        for i, page in enumerate(reader.pages)
    ]


def _extract_docx(path: Path) -> list[tuple[str, int | None, str | None]]:
    from docx import Document

    doc = Document(str(path))
    blocks = []
    current_section = None

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        if para.style and para.style.name and "Heading" in para.style.name:
            current_section = text

        blocks.append((text, None, current_section))

    # Include simple DOCX tables too. Business policies are often stored in
    # tables, and ignoring them can make an otherwise valid document appear
    # incomplete to RAG.
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            row_text = " | ".join(x for x in cells if x)
            if row_text:
                blocks.append((row_text, None, current_section))

    return blocks


def _extract_txt(path: Path) -> list[tuple[str, int | None, str | None]]:
    return [
        (
            path.read_text(encoding="utf-8", errors="replace"),
            None,
            None,
        )
    ]


def extract_document(
    path: Path,
    mime_type: str,
) -> list[tuple[str, int | None, str | None]]:
    if mime_type == "application/pdf":
        return _extract_pdf(path)

    if mime_type.endswith("wordprocessingml.document"):
        return _extract_docx(path)

    return _extract_txt(path)


def clean_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def chunk_blocks(
    blocks: Iterable[tuple[str, int | None, str | None]],
    max_chars: int = 1200,
    overlap: int = 180,
):
    chunks = []

    for raw, page, section in blocks:
        text = clean_text(raw)
        if not text:
            continue

        start = 0
        while start < len(text):
            end = min(len(text), start + max_chars)

            if end < len(text):
                boundary = text.rfind(" ", start, end)
                if boundary > start + 300:
                    end = boundary

            chunks.append((text[start:end], page, section))

            if end >= len(text):
                break

            start = max(0, end - overlap)

    return chunks


def delete_document_vectors(document: KnowledgeDocument):
    client = _load_dependencies()
    collection = client.get_or_create_collection(
        _collection_name(document.business_id)
    )

    ids = [c.vector_id for c in document.chunks]
    if ids:
        collection.delete(ids=ids)


def index_document(
    db: Session,
    document: KnowledgeDocument,
    path: Path,
):
    # Extraction and database chunk storage are the source of truth. Vector
    # indexing is an enhancement and must never make a valid upload fail.
    blocks = extract_document(path, document.mime_type)
    chunks = chunk_blocks(blocks)

    if not chunks:
        raise ValueError(
            "The document contains no extractable text. "
            "If this is a scanned PDF, use a text-based PDF/DOCX/TXT file."
        )

    ids = []
    texts = []
    metas = []

    for i, (text, page, section) in enumerate(chunks):
        vector_id = (
            f"doc{document.id}-v{document.version}-"
            f"{uuid.uuid4().hex}"
        )

        ids.append(vector_id)
        texts.append(text)
        metas.append(
            {
                "document_id": document.id,
                "business_id": document.business_id,
                "chunk_index": i,
                "page": page or 0,
                "section": section or "",
            }
        )

        db.add(
            KnowledgeChunk(
                document_id=document.id,
                chunk_index=i,
                page_number=page,
                section=section,
                content=text,
                vector_id=vector_id,
            )
        )

    document.chunk_count = len(chunks)
    document.status = "indexed"
    document.error_message = None

    # Commit the valid document/chunks before optional vector indexing. This
    # prevents Chroma or embedding infrastructure from turning a successful
    # document upload into a generic "Upload failed" response.
    db.commit()

    try:
        client = _load_dependencies()
        collection = client.get_or_create_collection(
            _collection_name(document.business_id),
            metadata={"hnsw:space": "cosine"},
        )
        vectors = _embed_texts(texts)
        collection.add(
            ids=ids,
            documents=texts,
            embeddings=vectors,
            metadatas=metas,
        )
    except Exception as exc:
        # Keep the searchable database chunks. Retrieval has a lexical
        # fallback, so RAG remains useful even when Chroma is unavailable.
        document.error_message = (
            f"Vector index unavailable: {type(exc).__name__}: {exc}"
        )[:2000]
        db.commit()


def _lexical_retrieve(
    db: Session,
    business_id: int,
    query: str,
    top_k: int,
) -> list[dict]:
    """Small dependency-free fallback for cloud environments."""
    terms = [
        x for x in re.findall(r"[a-z0-9]+", query.lower())
        if len(x) >= 2
    ]

    rows = (
        db.query(KnowledgeChunk, KnowledgeDocument)
        .join(
            KnowledgeDocument,
            KnowledgeDocument.id == KnowledgeChunk.document_id,
        )
        .filter(KnowledgeDocument.business_id == business_id)
        .all()
    )

    scored = []
    for chunk, doc in rows:
        haystack = chunk.content.lower()
        score = sum(haystack.count(term) for term in terms)
        if score:
            scored.append((score, chunk, doc))

    scored.sort(key=lambda x: x[0], reverse=True)

    return [
        {
            "document_id": doc.id,
            "document_name": doc.original_name,
            "page": chunk.page_number,
            "section": chunk.section,
            "content": chunk.content,
            "distance": 1.0 / (1.0 + score),
        }
        for score, chunk, doc in scored[:top_k]
    ]


def retrieve(
    db: Session,
    business_id: int,
    query: str,
    top_k: int = 5,
) -> list[dict]:
    try:
        client = _load_dependencies()
        collection = client.get_or_create_collection(
            _collection_name(business_id)
        )

        count = collection.count()
        if count:
            vector = _embed_texts([query])
            result = collection.query(
                query_embeddings=vector,
                n_results=min(top_k, count),
                where={"business_id": business_id},
                include=["documents", "metadatas", "distances"],
            )

            out = []
            for text, meta, distance in zip(
                result.get("documents", [[]])[0],
                result.get("metadatas", [[]])[0],
                result.get("distances", [[]])[0],
            ):
                doc = db.get(
                    KnowledgeDocument,
                    int(meta["document_id"]),
                )

                out.append(
                    {
                        "document_id": doc.id if doc else None,
                        "document_name": (
                            doc.original_name if doc else "Unknown"
                        ),
                        "page": meta.get("page") or None,
                        "section": meta.get("section") or None,
                        "content": text,
                        "distance": distance,
                    }
                )

            if out:
                return out
    except Exception:
        pass

    return _lexical_retrieve(
        db,
        business_id,
        query,
        top_k,
    )

